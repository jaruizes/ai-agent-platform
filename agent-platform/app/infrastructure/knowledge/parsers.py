from __future__ import annotations

import csv
import json
import mimetypes
import subprocess
import tempfile
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation

from app.domain.knowledge import ParsedDocument, ParsedSegment


class DocumentParser:
    SUPPORTED_EXTENSIONS = {
        ".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".md", ".json",
        ".doc", ".ppt", ".xls", ".odt", ".odp", ".ods",
    }

    def parse(self, path: str, *, name: str, mime_type: str | None = None) -> ParsedDocument:
        source = Path(path)
        suffix = source.suffix.lower()

        if suffix == ".pdf":
            return self._pdf(source, name)
        if suffix == ".docx":
            return self._docx(source, name)
        if suffix == ".pptx":
            return self._pptx(source, name)
        if suffix == ".xlsx":
            return self._xlsx(source, name)
        if suffix == ".csv":
            return self._csv(source, name)
        if suffix in {".txt", ".md"}:
            return ParsedDocument(
                title=name,
                segments=[ParsedSegment(source.read_text(encoding="utf-8", errors="replace"))],
                metadata={"format": suffix.lstrip(".")},
            )
        if suffix == ".json":
            data = json.loads(source.read_text(encoding="utf-8"))
            return ParsedDocument(
                title=name,
                segments=[ParsedSegment(json.dumps(data, ensure_ascii=False, indent=2))],
                metadata={"format": "json"},
            )
        if suffix in {".doc", ".ppt", ".xls", ".odt", ".odp", ".ods"}:
            return self._via_libreoffice_pdf(source, name)

        guessed = mime_type or mimetypes.guess_type(name)[0]
        raise ValueError(f"Unsupported document format: extension={suffix!r}, mimeType={guessed!r}")

    def _pdf(self, path: Path, name: str) -> ParsedDocument:
        doc = fitz.open(path)
        segments = [
            ParsedSegment(page.get_text("text").strip(), {"page": index + 1})
            for index, page in enumerate(doc)
            if page.get_text("text").strip()
        ]
        return ParsedDocument(
            title=name,
            segments=segments,
            metadata={"format": "pdf", "pageCount": len(doc)},
        )

    def _docx(self, path: Path, name: str) -> ParsedDocument:
        doc = DocxDocument(path)
        segments: list[ParsedSegment] = []
        for paragraph in doc.paragraphs:
            value = paragraph.text.strip()
            if value:
                segments.append(ParsedSegment(value, {"kind": "paragraph"}))
        for table_index, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            value = "\n".join(row for row in rows if row.strip())
            if value:
                segments.append(
                    ParsedSegment(value, {"kind": "table", "table": table_index + 1})
                )
        return ParsedDocument(title=name, segments=segments, metadata={"format": "docx"})

    def _pptx(self, path: Path, name: str) -> ParsedDocument:
        presentation = Presentation(path)
        segments: list[ParsedSegment] = []
        for slide_index, slide in enumerate(presentation.slides):
            parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = shape.text.strip()
                    if value:
                        parts.append(value)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            if parts:
                segments.append(
                    ParsedSegment(
                        "\n".join(parts),
                        {"slide": slide_index + 1, "kind": "slide"},
                    )
                )
        return ParsedDocument(
            title=name,
            segments=segments,
            metadata={"format": "pptx", "slideCount": len(presentation.slides)},
        )

    def _xlsx(self, path: Path, name: str) -> ParsedDocument:
        workbook = load_workbook(path, read_only=True, data_only=True)
        segments: list[ParsedSegment] = []
        for sheet in workbook.worksheets:
            lines: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    lines.append(" | ".join(values))
            if lines:
                segments.append(
                    ParsedSegment(
                        "\n".join(lines),
                        {"sheet": sheet.title, "kind": "sheet"},
                    )
                )
        return ParsedDocument(
            title=name,
            segments=segments,
            metadata={"format": "xlsx", "sheetCount": len(workbook.sheetnames)},
        )

    def _csv(self, path: Path, name: str) -> ParsedDocument:
        lines: list[str] = []
        with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
            reader = csv.reader(handle)
            for row in reader:
                lines.append(" | ".join(row))
        return ParsedDocument(
            title=name,
            segments=[ParsedSegment("\n".join(lines), {"kind": "table"})],
            metadata={"format": "csv"},
        )

    def _via_libreoffice_pdf(self, path: Path, name: str) -> ParsedDocument:
        with tempfile.TemporaryDirectory(prefix="knowledge-lo-") as tmp:
            completed = subprocess.run(
                [
                    "libreoffice",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    tmp,
                    str(path),
                ],
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
            output = Path(tmp) / f"{path.stem}.pdf"
            if completed.returncode != 0 or not output.exists():
                raise RuntimeError(
                    "LibreOffice conversion failed: "
                    f"{(completed.stderr or completed.stdout)[-2000:]}"
                )
            parsed = self._pdf(output, name)
            return ParsedDocument(
                title=parsed.title,
                segments=parsed.segments,
                metadata={**parsed.metadata, "convertedBy": "libreoffice", "sourceFormat": path.suffix.lower()},
            )
